"""Extract DOSM JMI Table 2 quarterly MAJOR occupational group offerings.

Use one DOSM workbook at a time. Prefer the latest reviewed release when
building a cumulative series; earlier releases may revise older figures.

Example (run from project root):
  python data_pipeline/clean_jmi.py "data/external/dosm/Data JMI Q2 2026.xlsx" --source-title "DOSM JMI Q2 2026" --source-url "https://www.dosm.gov.my/portal-main/release-content/big-data-analytics-job-market-insights-q22026" --output "data/processed/jmi_major_group_review.csv"

Output is deliberately separate from the existing validated dataset.
"""

import argparse
import csv
import re
from pathlib import Path
from openpyxl import load_workbook

FIELDS = ("period", "occupation_group", "job_offerings", "geography", "metric", "source_title", "source_url")
GROUPS = (
    "Managers", "Professionals", "Technicians and associate professionals",
    "Clerical support workers", "Service and sales workers",
    "Skilled agricultural, forestry, livestock and fishery workers",
    "Craft and related trades workers", "Plant and machine operators and assemblers",
    "Elementary Occupations",
)


def value_int(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)) and float(v).is_integer() and v >= 0:
        return int(v)
    if isinstance(v, str):
        v = v.replace(",", "").strip()
        if re.fullmatch(r"\d+", v):
            return int(v)
    return None


def find_table(sheet):
    """Find the labelled Occupation/Q1 header and annual-year row above it."""
    for row in sheet.iter_rows(min_row=1, max_row=100, values_only=False):
        for cell in row[:5]:
            if isinstance(cell.value, str) and cell.value.strip().lower() == "occupation":
                header_row = cell.row
                year_row = header_row - 1
                if year_row < 1:
                    continue
                group_col = cell.column
                quarter_columns = {}
                current_year = None
                # Forward-filled merged year labels in the quarterly column region.
                for col in range(group_col + 1, sheet.max_column + 1):
                    year_value = sheet.cell(year_row, col).value
                    if isinstance(year_value, (int, float)) and 2019 <= year_value <= 2100:
                        current_year = int(year_value)
                    quarter = sheet.cell(header_row, col).value
                    if isinstance(quarter, str) and re.fullmatch(r"Q[1-4]", quarter.strip(), re.I):
                        if current_year is None:
                            raise ValueError("Quarter header has no preceding year")
                        key = f"{current_year}-Q{quarter.strip()[-1]}"
                        if key in quarter_columns:
                            raise ValueError(f"Duplicate quarter column: {key}")
                        quarter_columns[key] = col
                if quarter_columns:
                    return header_row, group_col, quarter_columns
    raise ValueError("Could not find Table 2 'Occupation' quarter header")


def extract(path, source_title, source_url, start_year=2025, through_period=None):
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = next((s for s in workbook if "summary table" in s.title.lower()), None)
        if sheet is None:
            raise ValueError("No Summary Table worksheet detected")
        header, group_col, quarter_cols = find_table(sheet)
        selected = {p: c for p, c in quarter_cols.items()
                    if int(p[:4]) >= start_year and (through_period is None or p <= through_period)}
        if not selected:
            raise ValueError("No matching quarter columns")
        total = None
        groups = {}
        for rowno in range(header + 1, min(sheet.max_row, header + 30) + 1):
            label = sheet.cell(rowno, group_col).value
            if not isinstance(label, str):
                continue
            label = label.strip()
            if label.lower() in ("jumlah", "total"):
                total = {p: value_int(sheet.cell(rowno, c).value) for p, c in selected.items()}
            if label in GROUPS:
                groups[label] = {p: value_int(sheet.cell(rowno, c).value) for p, c in selected.items()}
            if label.lower().startswith("source:"):
                break
        if not groups or total is None:
            raise ValueError("Could not locate occupation groups and total row")
        records = []
        missing = []
        warnings = []
        for p in sorted(selected):
            if total[p] in (None, 0):
                # In older JMI workbooks, unreleased future quarters are often zero.
                continue
            s = 0
            for name in GROUPS:
                v = groups.get(name, {}).get(p)
                if v is None:
                    # Absence is not treated as an observed zero.
                    missing.append((p, name))
                    continue
                records.append({"period": p, "occupation_group": name, "job_offerings": v,
                                "geography": "Malaysia", "metric": "online_job_offerings_major_group",
                                "source_title": source_title, "source_url": source_url})
                s += v
            if not any(x[0] == p for x in missing) and s != total[p]:
                warnings.append(f"{p}: groups sum to {s:,}, published total {total[p]:,} (difference {total[p]-s:+,})")
        return records, missing, warnings
    finally:
        workbook.close()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("workbook", type=Path, help="Official DOSM JMI workbook")
    p.add_argument("--output", type=Path, default=Path("data/processed/jmi_major_group_review.csv"))
    p.add_argument("--source-title", required=True, help="Official release title e.g. DOSM JMI Q2 2026")
    p.add_argument("--source-url", required=True, help="Official release URL")
    p.add_argument("--start-year", type=int, default=2025)
    p.add_argument("--through-period", default=None, help="e.g. 2026-Q2")
    args = p.parse_args()
    if not args.workbook.is_file():
        p.error(f"Workbook not found: {args.workbook}")
    if args.output.resolve() == args.workbook.resolve():
        p.error("Output cannot overwrite the original workbook")
    try:
        if not args.source_url.startswith("https://"):
            raise ValueError("Official source URL must use HTTPS")
        records, missing, warnings = extract(args.workbook, args.source_title, args.source_url, args.start_year, args.through_period)
        if not records:
            raise ValueError("No released quarterly observations found")
    except (OSError, ValueError) as exc:
        p.error(str(exc))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(records)
    print(f"Wrote {len(records)} observations to {args.output}")
    for item in missing:
        print(f"MISSING: {item[0]} {item[1]} (not treated as zero)")
    for warning in warnings:
        print(f"WARNING: {warning}")
    print("REVIEW REQUIRED: source metadata and workbook revisions before integration")

if __name__ == "__main__":
    main()
